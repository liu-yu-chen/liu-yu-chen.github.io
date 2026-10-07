"""Read project-source outlines, tables, charts and text without publishing source files."""
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import re
import json
import pdfplumber
from docx import Document

root=Path(__file__).resolve().parents[1]
sources={
  'lifestyle':Path(r'D:\cancer_survey\Lifestyle_Network_Depressive_Symptoms_Manuscript.docx'),
  'abm':Path(r'C:\Users\Administrator\OneDrive\文档\xwechat_files\wxid_62ydzpu6yebf22_8853\msg\file\2026-09\agent_based_modeling_review_revised_v4.docx'),
  'report':Path(r'C:\Users\Administrator\OneDrive - University of Macau\hsci 7001 final report\Report.pdf'),
  'ev':Path(r'D:\ONE DRIVE_PERSONAL\OneDrive\文档\厦门市新能源汽车市场调研.docx'),
  'cisc':Path(r'C:\Users\Administrator\Downloads\FULL_REPORT_FOR_CISC7201.docx'),
}
out=root/'tmp/source_extracts';out.mkdir(parents=True,exist_ok=True)
for key,path in sources.items():
    if path.suffix.lower()=='.docx':
        doc=Document(path)
        lines=[]; headings=[]
        for para in doc.paragraphs:
            text=re.sub(r'\s+',' ',para.text).strip()
            if text:
                lines.append(text)
                if para.style and para.style.name.startswith('Heading'):headings.append(text)
        tables=[]
        for ti,table in enumerate(doc.tables):
            rows=[[re.sub(r'\s+',' ',cell.text).strip() for cell in row.cells] for row in table.rows]
            tables.append({'index':ti+1,'rows':len(rows),'cols':max(map(len,rows),default=0),'sample':rows[:9]})
        with ZipFile(path) as z:
            files=z.namelist()
            chart_files=[n for n in files if n.startswith('word/charts/chart') and n.endswith('.xml')]
            images=[{'name':n,'bytes':z.getinfo(n).file_size} for n in files if n.startswith('word/media/') and not n.endswith('/')]
        info={'file':path.name,'size':path.stat().st_size,'paragraphs':len(lines),'headings':headings[:100],'tables':tables[:18],'charts':chart_files[:25],'images':images[:40],'text':lines}
        (out/f'{key}.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8')
        print(f'[{key}] {path.name}: {len(lines)} nonempty paragraphs; {len(headings)} headings; {len(tables)} tables; {len(chart_files)} native charts; {len(images)} embedded images')
        print('HEADINGS: '+' | '.join(headings[:18]))
        if tables:
            print('TABLE EXAMPLES: '+json.dumps(tables[:2],ensure_ascii=False)[:1300])
    else:
        with pdfplumber.open(path) as pdf:
            pages=[page.extract_text(layout=False) or '' for page in pdf.pages]
            text='\n\n'.join(f'[PAGE {i+1}]\n{page}' for i,page in enumerate(pages))
            (out/f'{key}.txt').write_text(text,encoding='utf-8')
            print(f'[{key}] {path.name}: {len(pages)} PDF pages; {len(text)} extracted characters; {sum(len(p.extract_tables() or []) for p in pdf.pages)} detected tables')
            for i,page in enumerate(pages):
                items=[line.strip() for line in page.splitlines() if line.strip()]
                print(f'PAGE {i+1}: '+' | '.join(items[:5]))
                if i==7:break
print(f'Extracts saved in {out}')
